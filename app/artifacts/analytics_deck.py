"""Business Analysis PowerPoint Presentation Generator using python-pptx and real database data."""

import os
from pathlib import Path
from datetime import datetime
from typing import Optional
from sqlalchemy.orm import Session
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

from app.config import settings
from app.services.analytics_service import AnalyticsService
from app.services.inventory_service import InventoryService
from app.artifacts.charts import (
    generate_sales_trend_chart,
    generate_top_products_chart,
    generate_low_stock_chart,
)


def create_header(slide, title_text: str, subtitle_text: str = ""):
    """Helper to add consistent branded slide headers."""
    # Header container
    tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(8.4), Inches(1.0))
    tf = tx_box.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.bold = True
    p.font.size = Pt(24)
    p.font.color.rgb = RGBColor(26, 54, 93)  # Dark Blue

    if subtitle_text:
        p2 = tf.add_paragraph()
        p2.text = subtitle_text
        p2.font.size = Pt(12)
        p2.font.color.rgb = RGBColor(113, 128, 150)


def generate_analysis_pptx(period: str, session: Session, output_path: Optional[str] = None) -> str:
    """
    Generate a genuine 6-slide PowerPoint business review deck using live DB metrics and charts.
    Returns the absolute path to the generated PPTX file.
    """
    analytics_svc = AnalyticsService(session)
    inventory_svc = InventoryService(session)

    # 1. Fetch live metrics
    summary = analytics_svc.get_sales_summary(period=period)
    trend = analytics_svc.get_daily_sales_trend(days=7)
    top_products = analytics_svc.get_top_selling_products(days=7, limit=5)
    low_stock = inventory_svc.get_low_stock()

    # Generate charts from real data
    trend_chart_path = generate_sales_trend_chart(trend)
    top_chart_path = generate_top_products_chart(top_products)
    stock_chart_path = generate_low_stock_chart(low_stock)

    prs = Presentation()
    prs.slide_width = Inches(10.0)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # --- SLIDE 1: Title Slide ---
    slide1 = prs.slides.add_slide(blank_layout)
    tx_box = slide1.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(8.0), Inches(3.0))
    tf = tx_box.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = f"{settings.DEFAULT_SHOP_NAME}"
    p.font.bold = True
    p.font.size = Pt(36)
    p.font.color.rgb = RGBColor(26, 54, 93)

    p2 = tf.add_paragraph()
    p2.text = f"Executive Business Review & Sales Analysis ({period.capitalize()})"
    p2.font.bold = True
    p2.font.size = Pt(20)
    p2.font.color.rgb = RGBColor(43, 108, 176)

    p3 = tf.add_paragraph()
    p3.text = f"Generated autonomously on {datetime.now().strftime('%d %B %Y, %I:%M %p')}"
    p3.font.size = Pt(12)
    p3.font.color.rgb = RGBColor(113, 128, 150)

    # --- SLIDE 2: Sales Overview ---
    slide2 = prs.slides.add_slide(blank_layout)
    create_header(slide2, "1. Sales & Revenue Overview", f"Aggregated performance for period: {period}")

    # Stat Cards
    metrics = [
        ("Total Sales (Grand Total)", f"₹{summary['total_sales']:,.2f}"),
        ("Total Invoices Generated", f"{summary['bill_count']}"),
        ("Average Ticket Size", f"₹{summary['average_ticket_size']:,.2f}"),
        ("Total GST Tax Collected", f"₹{summary['total_tax']:,.2f}"),
    ]

    for idx, (label, val) in enumerate(metrics):
        x = Inches(0.8 + (idx % 2) * 4.3)
        y = Inches(1.8 + (idx // 2) * 2.2)
        box = slide2.shapes.add_textbox(x, y, Inches(4.0), Inches(1.8))
        tf = box.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = label
        p.font.size = Pt(13)
        p.font.color.rgb = RGBColor(113, 128, 150)

        p2 = tf.add_paragraph()
        p2.text = val
        p2.font.bold = True
        p2.font.size = Pt(28)
        p2.font.color.rgb = RGBColor(43, 108, 176)

    # --- SLIDE 3: Sales Trend ---
    slide3 = prs.slides.add_slide(blank_layout)
    create_header(slide3, "2. Daily Sales Trend", "7-day rolling revenue progression")
    if os.path.exists(trend_chart_path):
        slide3.shapes.add_picture(trend_chart_path, Inches(1.0), Inches(1.8), width=Inches(8.0))

    # --- SLIDE 4: Top Selling Products ---
    slide4 = prs.slides.add_slide(blank_layout)
    create_header(slide4, "3. Top Selling Products", "Best performing inventory by sales volume")
    if os.path.exists(top_chart_path):
        slide4.shapes.add_picture(top_chart_path, Inches(1.0), Inches(1.8), width=Inches(8.0))

    # --- SLIDE 5: Inventory & Low Stock ---
    slide5 = prs.slides.add_slide(blank_layout)
    create_header(slide5, "4. Inventory & Reorder Warnings", "Items requiring immediate purchase order")
    if os.path.exists(stock_chart_path):
        slide5.shapes.add_picture(stock_chart_path, Inches(1.0), Inches(1.8), width=Inches(8.0))

    # --- SLIDE 6: GST Breakdown & Key Insights ---
    slide6 = prs.slides.add_slide(blank_layout)
    create_header(slide6, "5. GST Compliance & Strategic Insights", "Statutory tax reconciliation & operational recommendations")

    tx_box = slide6.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(8.4), Inches(5.0))
    tf = tx_box.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "GST Tax Breakdown"
    p.font.bold = True
    p.font.size = Pt(16)
    p.font.color.rgb = RGBColor(26, 54, 93)

    p_gst1 = tf.add_paragraph()
    p_gst1.text = f"• Total Taxable Turnover: ₹{summary['total_subtotal']:,.2f}"
    p_gst1.font.size = Pt(13)

    p_gst2 = tf.add_paragraph()
    p_gst2.text = f"• Central GST (CGST): ₹{summary['total_cgst']:,.2f}"
    p_gst2.font.size = Pt(13)

    p_gst3 = tf.add_paragraph()
    p_gst3.text = f"• State GST (SGST): ₹{summary['total_sgst']:,.2f}"
    p_gst3.font.size = Pt(13)

    p_gst4 = tf.add_paragraph()
    p_gst4.text = f"• Combined GST Paid/Payable: ₹{summary['total_tax']:,.2f}"
    p_gst4.font.size = Pt(13)
    p_gst4.font.bold = True

    p_gap = tf.add_paragraph()
    p_gap.text = ""

    p_ins = tf.add_paragraph()
    p_ins.text = "Actionable Recommendations"
    p_ins.font.bold = True
    p_ins.font.size = Pt(16)
    p_ins.font.color.rgb = RGBColor(26, 54, 93)

    low_count = len(low_stock)
    p_rec1 = tf.add_paragraph()
    p_rec1.text = f"1. Stock Replenishment: {low_count} product(s) are below safety reorder threshold."
    p_rec1.font.size = Pt(13)

    p_rec2 = tf.add_paragraph()
    p_rec2.text = f"2. Payment Modes: Primary collections handled via {summary['payment_modes_breakdown']}."
    p_rec2.font.size = Pt(13)

    p_rec3 = tf.add_paragraph()
    p_rec3.text = "3. Credit Discipline: Ensure weekly customer Khata reminders to keep debt aging low."
    p_rec3.font.size = Pt(13)

    # Save output
    artifacts_dir = settings.ensure_artifacts_dir()
    if not output_path:
        filename = f"Sales_Analysis_{period}_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.pptx"
        output_path = str(artifacts_dir / filename)

    prs.save(output_path)
    return output_path
